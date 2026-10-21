"""Flatten a report result into one table and stream it as CSV or XLSX.

:func:`build_table` flattens the :class:`~hyperadmin_reports.grid.ReportGrid`
the HTML page renders, totals included, so the page, the CSV and the XLSX never
disagree. Both streamers are plain (sync) iterators; a ``StreamingResponse``
runs them in a worker thread, off the event loop.

``stream_xlsx`` uses an ``openpyxl`` *write-only* workbook: rows go to the
sheet's temporary file as they are appended instead of building a cell object
graph, and the finished zip is spooled to disk past a few MiB, so a million-cell
export stays flat in memory.
"""

from __future__ import annotations

import csv
import io
import re
import tempfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook

from hyperadmin_reports.grid import TOTAL_LABEL, ReportGrid, build_grid

if TYPE_CHECKING:
    from collections.abc import Iterator, Mapping

    from hyperadmin_reports.view import ReportResult, ReportView

#: Media type of an ``.xlsx`` workbook.
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
#: Title of the XLSX cover sheet listing the filters the export was run with.
FILTERS_SHEET_TITLE = "Filters used"

_CSV_ROWS_PER_CHUNK = 500
_XLSX_CHUNK_BYTES = 64 * 1024
_XLSX_SPOOL_BYTES = 8 * 1024 * 1024
_SHEET_TITLE_MAX = 31
_SHEET_TITLE_INVALID = re.compile(r"[\[\]:*?/\\]")


@dataclass(frozen=True)
class ReportTable:
    """A report flattened to a header and dense rows (``None`` marks an empty cell)."""

    header: list[str]
    rows: list[list[Any]]


def flatten_grid(grid: ReportGrid) -> ReportTable:
    """``grid`` as one header and dense rows, with its totals column(s) and row.

    The leading columns are the row dimensions; then one column per column key
    and aggregate (the key joined by ``" / "``, followed by the aggregate label
    when the report has several aggregates or no column dimension); then one
    "Total" column per totalled aggregate. A last "Total" row carries the column
    totals and the grand total(s) when anything can be totalled.
    """
    header = [*grid.row_labels, *grid.column_labels, *grid.total_labels]
    rows = [
        [*row.keys, *(c.value for c in row.cells), *(t.value for t in row.totals)]
        for row in grid.rows
    ]
    if grid.has_col_totals:
        blanks = [None] * (len(grid.row_labels) - 1)
        rows.append(
            [
                TOTAL_LABEL,
                *blanks,
                *(c.value for c in grid.col_totals),
                *(t.value for t in grid.grand_totals),
            ]
        )
    return ReportTable(header=header, rows=rows)


def build_table(report: type[ReportView], result: ReportResult) -> ReportTable:
    """Pivot every aggregate of ``result`` into one ``rows x (cols x values)`` table."""
    return flatten_grid(build_grid(report, result))


def stream_csv(table: ReportTable) -> Iterator[str]:
    """Yield ``table`` as UTF-8 CSV text (no BOM), header first, in chunks of rows."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    def flush() -> str:
        text = buffer.getvalue()
        buffer.seek(0)
        buffer.truncate()
        return text

    writer.writerow(table.header)
    for index, row in enumerate(table.rows, start=1):
        writer.writerow(["" if value is None else value for value in row])
        if index % _CSV_ROWS_PER_CHUNK == 0:
            yield flush()
    tail = flush()
    if tail:
        yield tail


def _xlsx_value(value: Any) -> Any:
    """A value openpyxl can store: numbers and naive dates as is, anything else as text."""
    if value is None or isinstance(value, (bool, int, float, Decimal, str)):
        return value
    if isinstance(value, datetime):
        # Excel has no time zones: store aware values as naive UTC.
        if value.tzinfo is not None:
            return value.astimezone(timezone.utc).replace(tzinfo=None)
        return value
    if isinstance(value, date):
        return value
    return str(value)


def sheet_title(title: str) -> str:
    """``title`` made a valid, distinct worksheet title (<= 31 chars, no ``[]:*?/\\``)."""
    cleaned = _SHEET_TITLE_INVALID.sub(" ", title).strip()[:_SHEET_TITLE_MAX].strip()
    if not cleaned or cleaned.lower() == FILTERS_SHEET_TITLE.lower():
        return "Report"
    return cleaned


def stream_xlsx(
    table: ReportTable,
    *,
    title: str,
    filters: Mapping[str, Any] | None = None,
    generated_at: datetime | None = None,
) -> Iterator[bytes]:
    """Yield ``table`` as an ``.xlsx`` workbook, in chunks.

    The first sheet, ``"Filters used"``, names the report, the (UTC) generation
    time and every filter the export was run with (``None`` when unfiltered); the
    second sheet, titled after the report, holds ``table``.
    """
    stamp = (generated_at or datetime.now(timezone.utc)).astimezone(timezone.utc)
    workbook = Workbook(write_only=True)
    cover = workbook.create_sheet(FILTERS_SHEET_TITLE)
    cover.append(["Report", title])
    cover.append(["Generated at (UTC)", stamp.replace(tzinfo=None)])
    cover.append([])
    cover.append(["Filter", "Value"])
    for name, value in (filters or {}).items():
        cover.append([name, _xlsx_value(value)])
    if not filters:
        cover.append(["None"])

    data = workbook.create_sheet(sheet_title(title))
    data.append(table.header)
    for row in table.rows:
        data.append([_xlsx_value(value) for value in row])

    with tempfile.SpooledTemporaryFile(max_size=_XLSX_SPOOL_BYTES) as spool:
        workbook.save(spool)
        spool.seek(0)
        while chunk := spool.read(_XLSX_CHUNK_BYTES):
            yield chunk
