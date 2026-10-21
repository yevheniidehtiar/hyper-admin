"""Unit tests for the report pages and CSV / XLSX export routes (st-v054-olap-03).

The first four tests map 1:1 to the scenarios of
``.meta/epics/epic-v054-olap-reporting/stories/featviews-report-routes-and-export.md``;
the rest cover the ReportPage design (bucket switch, totals, multi-level headers,
null vs zero, the 403 card, exports carrying the query string, the sidebar and
index), the registry and the plugin wiring. A real admin is mounted (the
``reports`` plugin is discovered from its entry point), in bridge mode where a
permission matters.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import anyio
import pytest
from fastapi import FastAPI, Header, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncEngine, async_sessionmaker, create_async_engine
from sqlmodel import Field, SQLModel
from sqlmodel.ext.asyncio.session import AsyncSession

from hyperadmin import Admin
from hyperadmin.auth.bridge import ExternalAuth
from hyperadmin.core.plugins import PluginRegistry
from hyperadmin.core.registry import site
from hyperadmin.core.settings import HyperAdminSettings
from hyperadmin_reports import (
    Count,
    Group,
    ReportView,
    Sum,
    TimeBucket,
    get_report,
    register_report,
    registered_reports,
    unregister_report,
)
from hyperadmin_reports.export import XLSX_MEDIA_TYPE
from hyperadmin_reports.plugin import ReportsPlugin
from hyperadmin_reports.routes import can_view, parse_period


class RrOrder(SQLModel, table=True):
    __tablename__ = "test_report_routes_order"

    id: int | None = Field(default=None, primary_key=True)
    sku: str
    amount: int
    created_at: datetime


class OrdersRevenue(ReportView):
    slug = "orders-revenue"
    description = "Revenue per SKU and month."
    model = RrOrder
    rows = [Group("sku")]
    cols = [TimeBucket("created_at", period="month")]
    values = [Sum("amount")]


class OrdersBySku(ReportView):
    slug = "orders-by-sku"
    label = "Orders by SKU"
    model = RrOrder
    rows = [Group("sku")]
    values = [Sum("amount")]


class OrdersBySkuAndAmount(ReportView):
    slug = "orders-by-sku-amount"
    model = RrOrder
    rows = [Group("sku"), Group("amount")]
    values = [Count()]


REPORTS = (OrdersRevenue, OrdersBySku, OrdersBySkuAndAmount)
GRID_HEADER = ["Sku", "2025-01", "2025-02", "Total"]
GRID_ROWS = [["A", 10, 20, 30], ["B", 5, None, 5], ["Total", 15, 20, 35]]

_ALLOWED = {"Authorization": "Bearer allowed"}
_DENIED = {"Authorization": "Bearer denied"}
_HX = {"HX-Request": "true"}


def _get_user(authorization: str | None = Header(default=None)) -> Any:
    tokens = {"Bearer allowed": "allowed", "Bearer denied": "denied"}
    return (
        SimpleNamespace(id=tokens[authorization], name=tokens[authorization], is_staff=True)
        if authorization in tokens
        else None
    )


def _has_permission(user: Any, codename: str) -> bool:
    return user.name == "allowed" and codename == "view_report_orders_revenue"


@pytest.fixture(autouse=True)
def _registered() -> Iterator[None]:
    saved_models = dict(site._registry)
    site._registry = {}
    saved_reports = registered_reports()
    for report in saved_reports:
        unregister_report(report)
    for report in REPORTS:
        register_report(report)
    yield
    for report in registered_reports():
        unregister_report(report)
    for report in saved_reports:
        register_report(report)
    site._registry = saved_models


@pytest.fixture
def engine(tmp_path: Path) -> AsyncEngine:
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'reports.db'}")

    async def _seed() -> None:
        async with engine.begin() as conn:
            await conn.run_sync(SQLModel.metadata.create_all, tables=[RrOrder.__table__])  # type: ignore[attr-defined]
        async with AsyncSession(engine) as session:
            for sku, amount, month in (("A", 10, 1), ("A", 20, 2), ("B", 5, 1)):
                session.add(
                    RrOrder(
                        sku=sku,
                        amount=amount,
                        created_at=datetime(2025, month, 3, tzinfo=timezone.utc),
                    )
                )
            await session.commit()

    anyio.run(_seed)
    return engine


def _clear(engine: AsyncEngine) -> None:
    async def _run() -> None:
        async with AsyncSession(engine) as session:
            await session.exec(delete(RrOrder))  # type: ignore[call-overload]
            await session.commit()

    anyio.run(_run)


def _add(engine: AsyncEngine, sku: str, amount: int, month: int) -> None:
    async def _run() -> None:
        async with AsyncSession(engine) as session:
            session.add(
                RrOrder(
                    sku=sku, amount=amount, created_at=datetime(2025, month, 3, tzinfo=timezone.utc)
                )
            )
            await session.commit()

    anyio.run(_run)


def _client(engine: AsyncEngine, *, bridge: bool = False, **admin_kwargs: Any) -> TestClient:
    app = FastAPI()
    if bridge:
        admin_kwargs["auth"] = ExternalAuth(get_user=_get_user, has_permission=_has_permission)
    Admin(
        app,
        engine=engine,
        settings=HyperAdminSettings(secret_key="reports-test-secret"),
        **admin_kwargs,
    ).mount("/admin")
    return TestClient(app, follow_redirects=False)


def _between(text: str, start: str, end: str) -> str:
    return text.split(start, 1)[1].split(end, 1)[0]


def _link(text: str, testid: str) -> str:
    """The whole ``<a ...>`` tag carrying ``testid``, whatever the attribute order."""
    return next(tag for tag in text.split("<a ") if f'data-testid="{testid}"' in tag).split(">", 1)[
        0
    ]


def test_html_grid_renders_pivot(engine: AsyncEngine) -> None:
    """
    Scenario: HTML grid renders pivot
      When  GET /admin/reports/orders-revenue/
      Then  response is 200 and body has a report-grid region with the crosstab
      And   the grid holds the pivoted cells, row and column totals, and the export menu
    """
    response = _client(engine).get("/admin/reports/orders-revenue/")

    assert response.status_code == 200
    assert 'data-testid="report-grid"' in response.text
    assert '<th scope="colgroup" colspan="2" class="ha-xtab__start">2025</th>' in response.text
    assert '<abbr title="2025-01">' not in response.text  # month labels stay in full
    assert '<th scope="col" class="ha-num ha-xtab__start">2025-01</th>' in response.text
    assert '<th scope="col" class="ha-num">2025-02</th>' in response.text
    assert '<th scope="col" rowspan="2" class="ha-num ha-xtab__total">Total</th>' in response.text
    assert response.text.count('data-testid="report-row"') == 2
    assert 'data-testid="report-totals"' in response.text
    assert '<th scope="row" class="ha-sticky-first">A</th>' in response.text
    assert ">30</td>" in response.text  # row total of A
    assert ">35</td>" in response.text  # grand total
    assert '/admin/reports/orders-revenue/export.csv"' in response.text
    assert '/admin/reports/orders-revenue/export.xlsx"' in response.text
    assert 'data-testid="report-export-csv"' in response.text
    assert 'data-testid="report-export-xlsx"' in response.text
    assert 'data-testid="report-export"' in response.text
    assert "Uses the current period and filters" in response.text
    assert 'data-testid="report-meta"' in response.text
    assert "Created At by month" in response.text
    assert "Revenue per SKU and month." in response.text
    assert 'data-testid="report-key"' in response.text
    assert "— no rows in that month" in response.text
    assert "/static/plugins/reports/reports.css" in response.text


def test_csv_stream_sets_attachment_disposition(engine: AsyncEngine) -> None:
    """
    Scenario: CSV stream sets attachment disposition
      When  GET /admin/reports/orders-revenue/export.csv
      Then  Content-Type is "text/csv"
      And   Content-Disposition is "attachment; filename=orders-revenue.csv"
      And   the body is the grid as BOM-free UTF-8 CSV, totals included
    """
    response = _client(engine).get("/admin/reports/orders-revenue/export.csv")

    assert response.status_code == 200
    assert response.headers["content-type"].split(";")[0] == "text/csv"
    assert response.headers["content-disposition"] == "attachment; filename=orders-revenue.csv"
    assert not response.content.startswith(b"\xef\xbb\xbf")
    assert response.text.splitlines() == [
        "Sku,2025-01,2025-02,Total",
        "A,10,20,30",
        "B,5,,5",
        "Total,15,20,35",
    ]


def test_xlsx_stream_is_openable_by_openpyxl(engine: AsyncEngine) -> None:
    """
    Scenario: XLSX stream is openable by openpyxl
      When  GET /admin/reports/orders-revenue/export.xlsx
      Then  Content-Type is the XLSX media type
      And   loading the body with openpyxl.load_workbook yields a data sheet matching the grid
      And   a "Filters used" cover sheet before it
    """
    response = _client(engine).get("/admin/reports/orders-revenue/export.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == "attachment; filename=orders-revenue.xlsx"
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Filters used", "Orders revenue"]
    rows = [list(row) for row in workbook["Orders revenue"].iter_rows(values_only=True)]
    assert rows == [GRID_HEADER, *GRID_ROWS]
    cover = list(workbook["Filters used"].iter_rows(values_only=True))
    assert cover[0] == ("Report", "Orders revenue")
    assert ("None", None) in cover


@pytest.mark.parametrize("path", ["", "export.csv", "export.xlsx"])
def test_permission_codename_gates_report_access(engine: AsyncEngine, path: str) -> None:
    """
    Scenario: permission codename gates report access
      Given the user lacks "view_report_orders_revenue"
      When  GET /admin/reports/orders-revenue/ (or either export)
      Then  response is 403 with the plugin's card naming the permission to ask for
      And   a user holding it gets 200
    """
    client = _client(engine, bridge=True)
    url = f"/admin/reports/orders-revenue/{path}"

    denied = client.get(url, headers=_DENIED)

    assert denied.status_code == 403
    assert 'data-testid="report-forbidden"' in denied.text
    assert "Can view report Orders revenue" in denied.text
    assert ">view_report_orders_revenue</code>" in denied.text
    assert 'data-testid="report-forbidden-back"' in denied.text
    assert 'data-testid="error-page"' not in denied.text
    assert client.get(url, headers=_ALLOWED).status_code == 200


def test_unknown_report_is_404(engine: AsyncEngine) -> None:
    """
    Scenario: an unregistered slug is not found
      When  GET /admin/reports/nope/
      Then  response is 404
    """
    assert _client(engine).get("/admin/reports/nope/").status_code == 404


def test_index_lists_only_reports_the_user_may_view(engine: AsyncEngine) -> None:
    """
    Scenario: the reports index hides reports the user cannot open
      Given the user holds only "view_report_orders_revenue"
      When  GET /admin/reports/
      Then  the index card links "Orders revenue" with its description and summary
      And   "Orders by SKU" is absent; the sidebar has the index and the permitted report
    """
    client = _client(engine, bridge=True)

    allowed = client.get("/admin/reports/", headers=_ALLOWED)
    denied = client.get("/admin/reports/", headers=_DENIED)

    assert allowed.status_code == 200
    assert 'data-testid="report-index"' in allowed.text
    assert '/admin/reports/orders-revenue/" data-testid="report-link-orders-revenue"' in (
        allowed.text
    )
    cards = _between(allowed.text, 'data-testid="report-index"', "</ul>")
    assert "Revenue per SKU and month." in cards
    assert "Rr order · sum of amount" in cards
    assert "Orders by SKU" not in cards
    assert 'data-testid="sidebar-plugin-reports-reports-index"' in allowed.text
    assert 'data-testid="nav-report-orders-revenue"' in allowed.text
    assert ">Reports</h3>" in allowed.text
    assert 'data-testid="report-index-empty"' in denied.text
    assert 'data-testid="report-index"' not in denied.text


def test_sidebar_links_every_report_for_an_unrestricted_user(engine: AsyncEngine) -> None:
    """
    Scenario: the sidebar has one entry per report under "Reports"
      Given no auth is configured
      When  GET /admin/reports/orders-revenue/
      Then  each report has a sidebar link with the bar-chart icon
      And   the current report is marked aria-current
    """
    page = _client(engine).get("/admin/reports/orders-revenue/").text

    link = _between(page, 'data-testid="nav-report-orders-revenue"', "</a>")
    assert 'aria-current="page"' in link
    assert "<svg" in link and "M3 3v18h18" in link  # the bar-chart-3 icon
    assert "Orders revenue</span>" in link
    assert 'data-testid="nav-report-orders-by-sku"' in page
    assert 'data-testid="nav-report-orders-by-sku-amount"' in page


def test_report_without_column_dimension_uses_session_factory(engine: AsyncEngine) -> None:
    """
    Scenario: a host session factory serves report queries
      Given Admin(session_factory=...) and a report with no column dimension
      When  GET /admin/reports/orders-by-sku/export.csv
      Then  the CSV has one column per row dimension and aggregate, and a totals row
    """
    factory = async_sessionmaker(engine, expire_on_commit=False)

    response = _client(engine, session_factory=factory).get(
        "/admin/reports/orders-by-sku/export.csv"
    )

    assert response.text.splitlines() == ["Sku,Sum of Amount", "A,30", "B,5", "Total,35"]


def test_report_without_time_bucket_has_no_bucket_switch(engine: AsyncEngine) -> None:
    """
    Scenario: the bucket switch only shows for reports with a time bucket
      When  GET /admin/reports/orders-by-sku/
      Then  there is no "Group by" control and no totals column, only the totals row
    """
    page = _client(engine).get("/admin/reports/orders-by-sku/").text

    assert 'data-testid="report-period"' not in page
    assert "ha-xtab__total" not in page
    assert 'data-testid="report-totals"' in page
    assert "— no rows in that group" in page


def test_bucket_switch_links_every_period_and_marks_the_current_one(
    engine: AsyncEngine,
) -> None:
    """
    Scenario: bucket switch
      When  GET /admin/reports/orders-revenue/?period=quarter
      Then  the "Group by" control links Auto, Day, Week, Month, Quarter and Year
      And   Quarter is current; each link is upgraded with hx-get to swap the page
      And   the grid buckets by quarter with "Q1" abbreviated under a 2025 group
      And   the export links carry the query string
    """
    page = _client(engine).get("/admin/reports/orders-revenue/?period=quarter").text

    seg = _between(page, 'data-testid="report-period"', "</div>")
    for value in ("auto", "day", "week", "month", "quarter", "year"):
        assert f'href="/admin/reports/orders-revenue/?period={value}"' in seg
        assert f'data-testid="report-period-{value}"' in seg
    assert 'aria-current="true"' in _link(seg, "report-period-quarter")
    assert 'aria-current="true"' not in _link(seg, "report-period-month")
    assert 'hx-get="/admin/reports/orders-revenue/?period=quarter"' in seg
    assert 'hx-target="#report-page"' in seg
    assert 'hx-push-url="true"' in seg
    assert 'data-testid="report-period-note"' not in page  # no note unless Auto is chosen
    assert '<abbr title="2025-Q1">Q1</abbr><span class="ha-sr-only"> 2025</span>' in page
    assert "Created At by quarter" in page
    assert '/admin/reports/orders-revenue/export.csv?period=quarter"' in page
    assert '/admin/reports/orders-revenue/export.xlsx?period=quarter"' in page
    assert page.count('data-testid="report-row"') == 2


def test_auto_period_explains_what_it_chose(engine: AsyncEngine) -> None:
    """
    Scenario: Auto says which bucket it picked and why
      When  GET /admin/reports/orders-revenue/?period=auto (data spans 32 days)
      Then  the note says "Auto: week, for a 32-day range" and the switch describes itself by it
      And   week labels abbreviate to "W01" under the year
      And   with no rows the note says there was nothing to span
    """
    client = _client(engine)

    page = client.get("/admin/reports/orders-revenue/?period=auto").text

    assert 'data-testid="report-period-note"' in page
    assert "Auto: week, for a 32-day range" in page
    assert 'aria-describedby="report-period-note"' in page
    assert '<abbr title="2025-W01">W01</abbr>' in page

    _clear(engine)
    empty = client.get("/admin/reports/orders-revenue/?period=auto").text

    assert "Auto: year, no dated rows to span" in empty
    assert 'data-testid="report-no-data"' in empty


def test_unknown_period_is_ignored(engine: AsyncEngine) -> None:
    """
    Scenario: an unknown ?period= falls back to the declared period
      When  GET /admin/reports/orders-revenue/?period=decade
      Then  the grid buckets by month and Month is current
    """
    page = _client(engine).get("/admin/reports/orders-revenue/?period=decade").text

    assert "Created At by month" in page
    assert 'aria-current="true"' in _link(page, "report-period-month")
    assert parse_period("decade") is None
    assert parse_period(None) is None
    assert parse_period("day") == "day"


def test_exports_honour_the_period_override(engine: AsyncEngine) -> None:
    """
    Scenario: exports use the bucket on screen
      When  GET export.csv and export.xlsx with ?period=year
      Then  both hold a single 2025 column, and the XLSX cover sheet lists the period
    """
    client = _client(engine)

    csv = client.get("/admin/reports/orders-revenue/export.csv?period=year")
    xlsx = client.get("/admin/reports/orders-revenue/export.xlsx?period=year")

    assert csv.text.splitlines() == ["Sku,2025,Total", "A,30,30", "B,5,5", "Total,35,35"]
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert ("period", "year") in list(workbook["Filters used"].iter_rows(values_only=True))


def test_htmx_request_gets_the_page_fragment(engine: AsyncEngine) -> None:
    """
    Scenario: the bucket switch swaps the report in place
      When  GET /admin/reports/orders-revenue/?period=year with HX-Request
      Then  the response is the #report-page fragment alone, without the document shell
    """
    response = _client(engine).get("/admin/reports/orders-revenue/?period=year", headers=_HX)

    assert response.status_code == 200
    assert response.text.lstrip().startswith('<div id="report-page"')
    assert "<html" not in response.text
    assert 'data-testid="sidebar"' not in response.text
    assert ">2025</th>" in response.text


def test_null_and_zero_cells_differ(engine: AsyncEngine) -> None:
    """
    Scenario: null vs zero
      Given sku C has orders adding up to zero in January and none in February
      When  GET /admin/reports/orders-revenue/
      Then  the January cell prints 0 as a value and February prints "—" with "no data"
    """
    _add(engine, "C", 0, 1)

    page = _client(engine).get("/admin/reports/orders-revenue/").text

    row = _between(page, '<th scope="row" class="ha-sticky-first">C</th>', "</tr>")
    cells = row.split("<td")[1:]
    assert "ha-xtab__zero" in cells[0] and ">0</td>" in cells[0]
    assert "&mdash;" in cells[1] and "> no data</span>" in cells[1]
    assert "ha-xtab__zero" not in cells[1]
    assert "ha-xtab__total" in cells[2] and ">0</td>" in cells[2]


def test_two_row_dimensions_render_as_row_groups(engine: AsyncEngine) -> None:
    """
    Scenario: row groups are quiet bands
      Given a report grouped by sku, then amount
      When  GET /admin/reports/orders-by-sku-amount/
      Then  each sku is a rowgroup header spanning the grid and its amounts are the rows
    """
    page = _client(engine).get("/admin/reports/orders-by-sku-amount/").text

    assert page.count('data-testid="report-rowgroup"') == 2
    assert '<th scope="rowgroup" colspan="2" class="ha-sticky-first">A</th>' in page
    assert '<th scope="row" class="ha-sticky-first">10</th>' in page
    assert '<th scope="row" class="ha-sticky-first">20</th>' in page
    assert "Sku / Amount</th>" in page
    assert ">3</td>" in page  # column total of the count


def test_empty_report_renders_no_data(engine: AsyncEngine) -> None:
    """
    Scenario: an empty result set renders a "No data" cell
      Given the report's table has no rows
      When  GET /admin/reports/orders-revenue/ with and without a query string
      Then  the grid shows one full-width "No data" cell naming the period
      And   with a query string it offers "Clear filters"; there is no totals row
    """
    _clear(engine)
    client = _client(engine)

    plain = client.get("/admin/reports/orders-revenue/").text
    filtered = client.get("/admin/reports/orders-revenue/?period=month").text

    cell = _between(plain, 'data-testid="report-no-data">', "</td>")
    assert "No data for created at by month." in cell
    assert "Clear filters" not in cell
    assert 'colspan="2"' in _between(plain, "<tbody>", 'data-testid="report-no-data"')
    assert 'data-testid="report-totals"' not in plain
    assert 'href="/admin/reports/orders-revenue/">Clear filters</a>' in filtered


@pytest.mark.anyio
async def test_can_view_without_a_user_is_403() -> None:
    """
    Scenario: auth is configured but the request carries no user
      Given an admin with auth and a permission checker
      When  can_view runs for a request without request.state.user
      Then  it raises 403
    """
    admin = SimpleNamespace(auth=object(), auth_backend=None, permission_checker=object())
    request = SimpleNamespace(state=SimpleNamespace())

    with pytest.raises(HTTPException) as exc:
        await can_view(admin, request, OrdersRevenue)  # type: ignore[arg-type]

    assert exc.value.status_code == 403


def test_register_report_rejects_unsafe_and_taken_slugs() -> None:
    """
    Scenario: report slugs are URL-safe and unique
      Given "orders-revenue" is registered
      When  another class registers that slug, or a slug with a "/"
      Then  register_report raises ValueError
      And   re-registering the same class is a no-op
    """

    class Clash(OrdersRevenue):
        pass

    class Unsafe(OrdersRevenue):
        slug = "a/b"

    with pytest.raises(ValueError, match="already registered by OrdersRevenue"):
        register_report(Clash)
    with pytest.raises(ValueError, match="is not URL-safe"):
        register_report(Unsafe)
    assert register_report(OrdersRevenue) is OrdersRevenue
    assert get_report("orders-revenue") is OrdersRevenue

    unregister_report(Clash)  # not registered: no-op

    assert get_report("orders-revenue") is OrdersRevenue


def test_plugin_declares_one_permission_per_codename() -> None:
    """
    Scenario: the plugin declares each registered report's permission once
      Given two reports, one sharing a custom codename with another
      When  ReportsPlugin is registered
      Then  the page, the index nav entry, one gated nav entry per report
      And   one permission per distinct codename are contributed
    """

    class Shared(OrdersBySku):
        slug = "orders-shared"
        label = "Orders shared"
        permission = "view_report_orders_by_sku"

    register_report(Shared)
    try:
        registry = PluginRegistry()
        assert registry.register(ReportsPlugin(), admin=SimpleNamespace(plugins=registry))
    finally:
        unregister_report(Shared)

    assert [r.slug for r in registered_reports()] == [
        "orders-revenue",
        "orders-by-sku",
        "orders-by-sku-amount",
    ]
    assert [p.prefix for p in registry.pages] == ["/reports"]
    entries = registry.nav_entries
    assert [e.url_name for e in entries] == ["reports-index", *["report-detail"] * 4]
    assert [e.section for e in entries] == ["Reports"] * 5
    assert entries[0].permission is None
    assert [
        (e.label, e.url_params["slug"], e.permission, e.testid, e.icon) for e in entries[1:]
    ] == [
        (
            "Orders revenue",
            "orders-revenue",
            "view_report_orders_revenue",
            "nav-report-orders-revenue",
            "bar-chart-3",
        ),
        (
            "Orders by SKU",
            "orders-by-sku",
            "view_report_orders_by_sku",
            "nav-report-orders-by-sku",
            "bar-chart-3",
        ),
        (
            "Orders by sku amount",
            "orders-by-sku-amount",
            "view_report_orders_by_sku_amount",
            "nav-report-orders-by-sku-amount",
            "bar-chart-3",
        ),
        (
            "Orders shared",
            "orders-shared",
            "view_report_orders_by_sku",
            "nav-report-orders-shared",
            "bar-chart-3",
        ),
    ]
    assert [(p.codename, p.name) for p in registry.permissions] == [
        ("view_report_orders_revenue", "Can view report Orders revenue"),
        ("view_report_orders_by_sku", "Can view report Orders by SKU"),
        ("view_report_orders_by_sku_amount", "Can view report Orders by sku amount"),
    ]
