"""Unit tests for ``hyperadmin_reports.export`` (grid flattening, CSV and XLSX streamers)."""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlmodel import Field, SQLModel

from hyperadmin_reports import Count, Group, ReportResult, ReportView, Sum, TimeBucket
from hyperadmin_reports.export import (
    ReportTable,
    build_table,
    sheet_title,
    stream_csv,
    stream_xlsx,
)


class ExOrder(SQLModel, table=True):
    __tablename__ = "test_report_export_order"

    id: int | None = Field(default=None, primary_key=True)
    sku: str
    amount: int
    created_at: datetime


class TwoValues(ReportView):
    slug = "two-values"
    model = ExOrder
    rows = [Group("sku")]
    cols = [TimeBucket("created_at", period="month")]
    values = [Sum("amount"), Count()]


def _result(records: list[dict[str, Any]]) -> ReportResult:
    return ReportResult(
        records=records,
        rows=["sku"],
        cols=["created_at"],
        values=["sum_amount", "count_all"],
        periods={"created_at": "month"},
        fills={"sum_amount": None, "count_all": 0},
    )


def _xlsx_rows(chunks: Any, sheet: str) -> list[tuple[Any, ...]]:
    workbook = load_workbook(io.BytesIO(b"".join(chunks)))
    return list(workbook[sheet].iter_rows(values_only=True))


def test_build_table_labels_each_value_under_each_column_key() -> None:
    """
    Scenario: a report with several aggregates flattens to one column per (key, aggregate)
      Given a report with Sum and Count over a month bucket
      When  build_table runs on its result
      Then  each header is "<column key> / <aggregate label>" and missing cells take the fill
    """
    table = build_table(
        TwoValues,
        _result(
            [
                {"sku": "A", "created_at": "2025-01", "sum_amount": 10, "count_all": 1},
                {"sku": "B", "created_at": "2025-02", "sum_amount": 5, "count_all": 2},
            ]
        ),
    )

    assert table.header == [
        "Sku",
        "2025-01 / Sum of Amount",
        "2025-01 / Count of Rows",
        "2025-02 / Sum of Amount",
        "2025-02 / Count of Rows",
    ]
    assert table.rows == [["A", 10, 1, None, 0], ["B", None, 0, 5, 2]]


def test_build_table_of_an_empty_result_has_only_row_headers() -> None:
    """
    Scenario: an empty result set flattens to a header without rows
      When  build_table runs on a result with no records
      Then  the header is the row dimensions only and there are no rows
    """
    assert build_table(TwoValues, _result([])) == ReportTable(header=["Sku"], rows=[])


def test_stream_csv_yields_rows_in_chunks() -> None:
    """
    Scenario: a long CSV is streamed in chunks of rows
      Given a 1001-row table
      When  stream_csv runs
      Then  it yields three chunks whose concatenation is the whole CSV
    """
    table = ReportTable(header=["n"], rows=[[i] for i in range(1001)])

    chunks = list(stream_csv(table))

    assert len(chunks) == 3
    lines = "".join(chunks).splitlines()
    assert lines[0] == "n"
    assert lines[-1] == "1000"
    assert len(lines) == 1002


def test_stream_csv_of_an_exact_chunk_multiple_has_no_empty_tail() -> None:
    """
    Scenario: no empty chunk follows a full last chunk
      Given a 500-row table (one full chunk)
      When  stream_csv runs
      Then  it yields exactly one chunk
    """
    assert len(list(stream_csv(ReportTable(header=["n"], rows=[[i] for i in range(500)])))) == 1


def test_stream_xlsx_stores_typed_values_and_lists_filters() -> None:
    """
    Scenario: XLSX cells keep numbers and dates, and the cover sheet lists the filters
      Given cells of every supported type and a filter
      When  stream_xlsx runs
      Then  numbers and dates are typed cells, aware datetimes naive UTC, other values text
      And   the "Filters used" sheet lists the filter and the generation time
    """
    aware = datetime(2025, 1, 2, 3, 0, tzinfo=timezone(timedelta(hours=2)))
    table = ReportTable(
        header=["a", "b", "c", "d", "e", "f"],
        rows=[[Decimal("1.5"), aware, datetime(2025, 1, 2), date(2025, 1, 3), object, None]],
    )
    generated = datetime(2025, 6, 1, 12, 0, tzinfo=timezone.utc)

    chunks = list(
        stream_xlsx(table, title="Sales", filters={"region": "EU"}, generated_at=generated)
    )

    (header, row) = _xlsx_rows(chunks, "Sales")
    assert header == ("a", "b", "c", "d", "e", "f")
    assert row[0] == 1.5
    assert row[1] == datetime(2025, 1, 2, 1, 0)
    assert row[2] == datetime(2025, 1, 2)
    assert row[3] == datetime(2025, 1, 3)  # openpyxl reads dates back as datetimes
    assert row[4] == str(object)
    assert row[5] is None
    cover = _xlsx_rows(chunks, "Filters used")
    assert ("Generated at (UTC)", datetime(2025, 6, 1, 12, 0)) in cover
    assert ("region", "EU") in cover
    assert ("None", None) not in cover


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("Revenue: Q1/Q2", "Revenue  Q1 Q2"),
        ("x" * 40, "x" * 31),
        ("[]", "Report"),
        ("filters USED", "Report"),
    ],
)
def test_sheet_title_is_valid_and_distinct(title: str, expected: str) -> None:
    """
    Scenario: the data sheet title is a valid worksheet name
      When  sheet_title cleans a report label
      Then  invalid characters are replaced, it is cut to 31 characters
      And   an empty or cover-sheet-clashing title falls back to "Report"
    """
    assert sheet_title(title) == expected
